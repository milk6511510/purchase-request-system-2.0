from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUTPUT = Path("purchase_request_backend_template.xlsx").resolve()

SHEETS = {
    "請購主表": [
        "單號", "建立時間", "更新時間", "員工編號", "公司信箱", "申請人", "部門代號", "申請類型",
        "申請日期", "須用日期", "廠商名稱", "統一編號", "稅別", "合計金額", "簽核狀態", "目前關卡",
        "類別", "檢附資料", "用途說明", "驗收/其他說明", "照片數", "簽名狀態", "Notion頁面連結", "原始資料"
    ],
    "品項明細": [
        "單號", "項次", "品項名稱", "規格", "品項編號", "單位", "申購數量", "核准數量", "單價", "金額", "備註"
    ],
    "簽核紀錄": [
        "單號", "動作時間", "關卡", "動作", "簽核人工號", "簽核人", "簽核人信箱", "說明", "通知信箱", "Notion頁面連結"
    ],
    "使用者權限": [
        "員工編號", "姓名", "公司信箱", "部門代號", "權限", "啟用", "建立時間", "更新時間", "備註"
    ],
    "系統設定": [
        "設定項目", "設定值", "說明"
    ],
}

SEED_ROWS = {
    "請購主表": [[
        "REQ-EXAMPLE", "2026-07-23T10:00:00+08:00", "2026-07-23T10:00:00+08:00",
        "A001", "admin@example.com", "系統管理者", "ADMIN", "一般", "2026-07-23",
        "2026-07-30", "範例廠商", "12345678", "含稅", 1000, "待單位主管簽核",
        "單位主管", "雜項購置", "報價單", "範例用途說明", "", 1, "已簽名", "", "{}"
    ]],
    "品項明細": [["REQ-EXAMPLE", 1, "範例品項", "規格", "ITEM001", "PCS", 1, "", 1000, "=G2*I2", ""]],
    "簽核紀錄": [["REQ-EXAMPLE", "2026-07-23T10:00:00+08:00", "申請人", "已簽名", "A001", "系統管理者", "admin@example.com", "申請人送出", "admin@example.com", ""]],
    "使用者權限": [
        ["A001", "系統管理者", "", "ADMIN", "最高權限", True, "2026-07-23T10:00:00+08:00", "", "最高權限"],
        ["M001", "單位主管範例", "", "ADM01", "單位主管", True, "2026-07-23T10:00:00+08:00", "", ""],
        ["G001", "總經理範例", "", "GM", "總經理", True, "2026-07-23T10:00:00+08:00", "", ""],
    ],
    "系統設定": [
        ["可修改時限小時", "24", "使用者送出後可自行修改的小時數"],
        ["預設最高權限員工編號", "A001", "預設最高權限者"],
        ["資料表版本", "1.0", "請購單後台資料表版本"],
    ],
}

WIDTHS = {
    "單號": 18, "建立時間": 24, "更新時間": 24, "員工編號": 14, "公司信箱": 28,
    "申請人": 16, "部門代號": 16, "申請類型": 12, "申請日期": 14, "須用日期": 14,
    "廠商名稱": 22, "統一編號": 14, "稅別": 10, "合計金額": 14, "簽核狀態": 20,
    "目前關卡": 18, "類別": 22, "檢附資料": 22, "用途說明": 34, "驗收/其他說明": 30,
    "原始資料": 36, "品項名稱": 24, "備註": 26, "設定項目": 24, "設定值": 22, "說明": 42
}

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
subtle_fill = PatternFill("solid", fgColor="EAF3F8")
thin = Side(style="thin", color="D9E2EC")
border = Border(bottom=thin)

for sheet_name, headers in SHEETS.items():
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in SEED_ROWS.get(sheet_name, []):
        ws.append(row)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.fill = subtle_fill if cell.row == 2 else PatternFill(fill_type=None)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = WIDTHS.get(header, 16)

    ws.freeze_panes = "A2"
    table_ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    table = Table(displayName=f"Table{len(wb.worksheets)}", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.sheet_view.showGridLines = False

role_validation = DataValidation(type="list", formula1='"一般使用者,單位主管,總經理,最高權限"', allow_blank=False)
active_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
request_type_validation = DataValidation(type="list", formula1='"一般,急件"', allow_blank=True)
status_validation = DataValidation(type="list", formula1='"待單位主管簽核,待總經理簽核,已核准,已退回,已註銷"', allow_blank=True)

users = wb["使用者權限"]
users.add_data_validation(role_validation)
users.add_data_validation(active_validation)
role_validation.add("E2:E500")
active_validation.add("F2:F500")

requests = wb["請購主表"]
requests.add_data_validation(request_type_validation)
requests.add_data_validation(status_validation)
request_type_validation.add("H2:H1000")
status_validation.add("O2:O1000")

wb.save(OUTPUT)
print(OUTPUT)
